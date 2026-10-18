"""Sinh template Word mau (don mua hang) dung du cu phap cua ZCL_WORD_DOCX_TEST + anh chu ky mau.

    python tools/build_template.py template/demo_template.docx [template/demo_sign.png]

Cu phap dung trong template:
  ${tag}                       field thuong (xuong dong trong gia tri -> xuong dong trong Word)
  ${table:items.x}             dong bang lap theo mang items; ${rowif:isGroup} chon dong mau theo tung dong
  ${#each:copies} ... ${/each:copies}   lap ca khoi (moi lien 1 trang)
  ${#if:hasZeroPrice} ... ${/if:hasZeroPrice}   giu khoi khi field co gia tri; ${#ifnot:x} nguoc lai
  ${image:sign|3}              chen anh PNG/JPEG rong 3 cm
Co san 1 tag bi cat thanh nhieu run (${supplier|Name}) de thu buoc gop tag.
"""
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Pt, RGBColor

try:
    sys.stdout.reconfigure(encoding="utf-8")
except AttributeError:
    pass


def label(p, text, value):
    p.add_run(text).bold = True
    p.add_run(value)


def bold_cell(cell, text, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    p.add_run(text).bold = True


def main(out, sign_out=None):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)

    doc.sections[0].header.paragraphs[0].text = "Company Code: ${companyCode}"
    doc.sections[0].footer.paragraphs[0].text = "In bởi: ${printedBy}"

    doc.add_paragraph("${#each:copies}")

    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    c.add_run("${label}").italic = True

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("ĐƠN MUA HÀNG")
    r.bold = True
    r.font.size = Pt(16)

    label(doc.add_paragraph(), "Số đơn: ", "${purchaseOrder}")
    label(doc.add_paragraph(), "Ngày: ", "${poDate}")
    p = doc.add_paragraph()
    p.add_run("Nhà cung cấp: ").bold = True
    p.add_run("${supplier} - ${supplier")
    p.add_run("Name}").italic = True
    label(doc.add_paragraph(), "Tiền tệ: ", "${currency}")

    heads = ["Dòng", "Vật tư", "Mô tả", "SL", "ĐVT", "Đơn giá", "Thành tiền"]
    tags = ["item", "material", "text", "qty", "unit", "price", "amount"]
    tbl = doc.add_table(rows=5, cols=len(heads))
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(heads):
        bold_cell(tbl.rows[0].cells[i], h)
    # dong tieu de nhom
    g = tbl.rows[1].cells[0].merge(tbl.rows[1].cells[6])
    g.text = ""
    g.paragraphs[0].add_run("${rowif:isGroup}${table:items.groupLabel}").bold = True
    # dong hang
    for i, tg in enumerate(tags):
        tbl.rows[2].cells[i].text = "${table:items.%s}" % tg
    tbl.rows[2].cells[0].text = "${rowif:isItem}${table:items.item}"
    # dong cong nhom
    s = tbl.rows[3].cells[0].merge(tbl.rows[3].cells[5])
    s.text = ""
    s.paragraphs[0].add_run("${rowif:isSubtotal}${table:items.groupLabel}").italic = True
    tbl.rows[3].cells[6].text = "${table:items.amount}"
    # tong
    total = tbl.rows[4].cells[0].merge(tbl.rows[4].cells[5])
    bold_cell(total, "TỔNG CỘNG")
    tbl.rows[4].cells[6].text = "${total}"

    label(doc.add_paragraph(), "Bằng chữ: ", "${totalWords}")

    doc.add_paragraph("${#if:hasZeroPrice}")
    w = doc.add_paragraph().add_run("Lưu ý: đơn có dòng chưa có giá, cần bổ sung trước khi duyệt.")
    w.bold = True
    w.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    doc.add_paragraph("${/if:hasZeroPrice}")

    doc.add_paragraph("${#if:note}")
    label(doc.add_paragraph(), "Ghi chú: ", "${note}")
    doc.add_paragraph("${/if:note}")

    sign = doc.add_table(rows=3, cols=2)
    for i, txt in enumerate(["Người lập", "Người duyệt"]):
        bold_cell(sign.rows[0].cells[i], txt, WD_ALIGN_PARAGRAPH.CENTER)
    sign.rows[1].cells[0].text = "${image:sign|3}"
    sign.rows[1].cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    sign.rows[2].cells[0].text = "${printedBy}"
    sign.rows[2].cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("${#ifnot:isLast}")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_paragraph("${/if:isLast}")

    doc.add_paragraph("${/each:copies}")

    doc.save(out)
    print("da tao", out)

    if sign_out:
        from PIL import Image, ImageDraw
        img = Image.new("RGBA", (300, 120), (255, 255, 255, 0))  # nen trong suot, khong vien
        d = ImageDraw.Draw(img)
        d.line([(20, 90), (80, 30), (120, 85), (170, 25), (220, 80), (280, 40)], fill=(20, 40, 160, 255), width=5)
        img.save(sign_out)
        print("da tao", sign_out)


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "demo_template.docx", sys.argv[2] if len(sys.argv) > 2 else None)
